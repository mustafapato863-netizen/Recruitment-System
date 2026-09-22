from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

path = r"docs\handover\RecruitFlow_IT_Handover_Presentation.pptx"
prs = Presentation(path)
slide = prs.slides[0]

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xB8, 0xD4, 0xEB)

box = slide.shapes.add_textbox(Inches(0.8), Inches(4.85), Inches(11), Inches(0.95))
tf = box.text_frame
tf.word_wrap = True
p0 = tf.paragraphs[0]
r0 = p0.add_run()
r0.text = "Presented by"
r0.font.size = Pt(13)
r0.font.color.rgb = LIGHT
r0.font.name = "Calibri"

p1 = tf.add_paragraph()
r1 = p1.add_run()
r1.text = "Mustafa Zainhom"
r1.font.size = Pt(22)
r1.font.bold = True
r1.font.color.rgb = WHITE
r1.font.name = "Calibri"

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "HRIS Performance Specialist  ·  Saudi German Health"
r2.font.size = Pt(15)
r2.font.color.rgb = LIGHT
r2.font.name = "Calibri"

close = prs.slides[-1]
cbox = close.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11), Inches(0.5))
cr = cbox.text_frame.paragraphs[0].add_run()
cr.text = "Mustafa Zainhom  ·  HRIS Performance Specialist"
cr.font.size = Pt(16)
cr.font.bold = True
cr.font.color.rgb = WHITE
cr.font.name = "Calibri"

prs.save(path)
print("Updated", path)
