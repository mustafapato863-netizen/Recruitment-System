"""Generate RecruitFlow IT Department handover presentation."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml

# Brand colors (Saudi German Health / RecruitFlow)
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x00, 0x84, 0xCE)
GREEN = RGBColor(0x00, 0xA8, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
CARD = RGBColor(0xE8, 0xF4, 0xFC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def set_run(run, text, size=18, bold=False, color=SLATE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(slide, left, top, width, height, lines, default_size=18, default_color=SLATE, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, tuple):
            text, size, bold, color = line
        else:
            text, size, bold, color = line, default_size, (bold_first and i == 0), default_color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        set_run(run, text, size=size, bold=bold, color=color)
        p.space_after = Pt(6)
    return box


def add_title_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, W, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), W, Inches(0.08), TEAL)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.22), Inches(12), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    set_run(run, title, size=28, bold=True, color=WHITE)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.6), Inches(0.62), Inches(12), Inches(0.35))
        run2 = box2.text_frame.paragraphs[0].add_run()
        set_run(run2, subtitle, size=14, bold=False, color=RGBColor(0xB8, 0xD4, 0xEB))
    # footer
    add_rect(slide, 0, H - Inches(0.4), W, Inches(0.4), LIGHT)
    foot = slide.shapes.add_textbox(Inches(0.6), H - Inches(0.32), Inches(10), Inches(0.25))
    r = foot.text_frame.paragraphs[0].add_run()
    set_run(r, "RecruitFlow  |  IT Department Handover  |  Confidential", size=11, color=MUTED)
    page = slide.shapes.add_textbox(W - Inches(1.5), H - Inches(0.32), Inches(1.2), Inches(0.25))
    page.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        set_run(run, f"•  {item}", size=size, color=SLATE)
        p.space_after = Pt(8)
    return box


def card(slide, left, top, width, height, title, body_lines):
    add_rect(slide, left, top, width, height, CARD)
    add_rect(slide, left, top, Inches(0.08), height, TEAL)
    add_text(slide, left + Inches(0.25), top + Inches(0.15), width - Inches(0.35), Inches(0.35),
             [(title, 15, True, NAVY)])
    bullets(slide, left + Inches(0.2), top + Inches(0.5), width - Inches(0.35), height - Inches(0.6),
            body_lines, size=13)


# ========== SLIDES ==========

# 1 Title
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, Inches(5.6), W, Inches(1.9), TEAL)
add_text(s, Inches(0.8), Inches(2.0), Inches(11), Inches(1),
         [("RecruitFlow", 48, True, WHITE),
          ("System Handover Presentation", 28, False, RGBColor(0xB8, 0xD4, 0xEB))])
add_text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(1),
         [("Prepared for: IT Department", 18, False, WHITE),
          ("Saudi German Health  ·  Internal Recruitment Platform", 16, False, RGBColor(0xB8, 0xD4, 0xEB)),
          ("September 2026", 16, False, RGBColor(0xB8, 0xD4, 0xEB))])
add_text(s, Inches(0.8), Inches(6.0), Inches(11), Inches(0.8),
         [("Purpose: technical handover for ownership, operations, and support readiness", 15, False, WHITE)])

# 2 Agenda
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Agenda", "What we will cover today")
items = [
    "1. Product overview & business scope",
    "2. Architecture & repository layout",
    "3. Technology stack",
    "4. Data, storage & integrations",
    "5. Security, identity & access control",
    "6. Functional modules (end-to-end hiring lifecycle)",
    "7. Recent delivery (UX density, trust, fit/skills)",
    "8. Environments, deployment & runbooks",
    "9. Quality gates, known deferrals & risks",
    "10. Handover checklist & recommended next steps",
]
bullets(s, Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.2), items, size=18)

# 3 What is RecruitFlow
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "What is RecruitFlow?", "Internal recruitment & hiring operations platform")
bullets(s, Inches(0.8), Inches(1.4), Inches(6.2), Inches(5), [
    "Enterprise recruitment operations system for Saudi German Health",
    "Covers vacancy request → approval → posting → sourcing → interview → offer → hire",
    "Modular monolith: one monorepo, separately deployable web, API, and worker apps",
    "Business rules enforced server-side; PostgreSQL is system of record",
    "Not a full HCM/ERP replacement — focused recruitment module",
], size=16)
card(s, Inches(7.4), Inches(1.5), Inches(5.2), Inches(4.8), "In scope for IT ownership", [
    "Application hosting & runtime",
    "Database & migrations",
    "Secrets / environment config",
    "Access control & audit posture",
    "Backups, monitoring, support",
    "CI/CD and release process",
])

# 4 Architecture
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "High-Level Architecture", "Modular monolith with clear service boundaries")
cards = [
    (Inches(0.5), "apps/web", ["React + Vite UI", "Role-gated routes", "Design-system tokens"]),
    (Inches(3.6), "apps/api", ["NestJS REST API", "Auth, RBAC, workflows", "Prisma repositories"]),
    (Inches(6.7), "apps/worker", ["BullMQ jobs", "Email / async work", "No SoR for business state"]),
    (Inches(9.8), "Shared pkgs", ["contracts", "validation (Zod)", "config + design-system"]),
]
for left, title, body in cards:
    card(s, left, Inches(1.5), Inches(2.9), Inches(3.6), title, body)
add_text(s, Inches(0.6), Inches(5.4), Inches(12), Inches(1.2), [
    ("Data plane: PostgreSQL (authoritative)  ·  Redis/BullMQ (queues)  ·  S3-compatible / volume storage for documents", 14, False, MUTED),
    ("Dependency rule: API owns authorization & workflow rules; UI never becomes source of truth for permissions", 14, False, MUTED),
])

# 5 Tech stack
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Technology Stack", "Current production-oriented toolchain")
card(s, Inches(0.5), Inches(1.4), Inches(4), Inches(4.8), "Frontend", [
    "React 19 + TypeScript",
    "Vite 8",
    "React Router 7",
    "Tailwind CSS 4 + design tokens",
    "Vitest + Testing Library",
    "i18next (English-only scaffold)",
])
card(s, Inches(4.7), Inches(1.4), Inches(4), Inches(4.8), "Backend & Data", [
    "NestJS 11",
    "Prisma + PostgreSQL",
    "Passport JWT (access + refresh)",
    "Zod shared validation",
    "BullMQ worker process",
    "Affinda resume parsing integration",
])
card(s, Inches(8.9), Inches(1.4), Inches(4), Inches(4.8), "Tooling & Ops", [
    "pnpm workspaces (Node 24 / pnpm 11)",
    "Docker multi-stage (api/web/worker)",
    "compose.production.yml",
    "GitHub + GitLab remotes",
    "ESLint + typecheck + token gate",
    "Browser/Python test scripts",
])

# 6 Repo layout
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Repository Layout", "Single monorepo — independently buildable apps")
bullets(s, Inches(0.8), Inches(1.4), Inches(12), Inches(5.5), [
    "apps/web          — React + Vite SPA",
    "apps/api          — NestJS modular API",
    "apps/worker       — Background / scheduled jobs",
    "packages/contracts, validation, config, design-system — shared libraries",
    "database/         — Prisma schema, migrations, seeds, DB tests",
    "docs/             — Architecture, setup, design system, operations",
    "deploy/           — Runtime helpers, nginx template, Hostinger/Vercel notes",
    "Remotes: GitHub origin (Recruitment-System) and GitLab mirror",
], size=17)

# 7 Security
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Security & Access Control", "IT-critical controls")
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(5), "Identity & authorization", [
    "JWT access + refresh cookies (httpOnly, SameSite)",
    "Server-side permission checks on every sensitive route",
    "RBAC roles (recruiter, hiring manager, approver, HR admin, interviewer, admin)",
    "Frontend PermissionGate mirrors — API remains authoritative",
    "Audit log for sensitive actions",
])
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(5), "Hardening posture", [
    "Security headers (CSP, HSTS in prod, nosniff, frame options)",
    "CORS locked to configured WEB_ORIGIN",
    "Normalized safe API error envelopes",
    "Fail-fast environment validation on startup",
    "Design-token / lint / typecheck quality gates",
])

# 8 Modules lifecycle
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Functional Modules — Hiring Lifecycle", "What the business runs day-to-day")
flow = [
    ("1. Demand", ["Vacancy requests", "Approval inbox", "Headcount / org data"]),
    ("2. Attract", ["Vacancies / JD", "Public careers pages", "CV intake & import"]),
    ("3. Select", ["Candidates DB", "Smart sourcing & fit", "Compare & applications"]),
    ("4. Assess", ["Interviews / calendar", "Scorecards", "Tasks & SLA"]),
    ("5. Hire", ["Offers & approvals", "Final hiring approval", "Hiring cases / joining"]),
]
for i, (title, body) in enumerate(flow):
    left = Inches(0.4 + i * 2.55)
    card(s, left, Inches(1.5), Inches(2.4), Inches(4.6), title, body)

# 9 Key surfaces
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Key Operational Surfaces", "Primary screens IT should know for support")
bullets(s, Inches(0.8), Inches(1.4), Inches(12), Inches(5.5), [
    "Dashboards — employee / manager / recruitment command views",
    "Vacancy requests & vacancies workspace (including JD import)",
    "Candidates, CV Bank, Smart Sourcing & Match Engine",
    "Applications pipeline (list + kanban) with fit scorecards",
    "Interviews calendar & detail; candidate self-schedule links",
    "Offers, offer approval inbox, final hiring approval inbox",
    "Hiring / joining cases; reports & audit log",
    "Admin: users/roles, master data, settings, bulk import",
], size=17)

# 10 Recent delivery
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Recent Delivery (Sep 2026)", "Merged to main — relevant for support & regression")
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(5), "UX density & trust (Phases F–I)", [
    "Denser shell (brand/header/user chip)",
    "Unified page titles & compact KPI cards",
    "Responsive Wave-1 padding / actions",
    "Confirm + toast trust on Offer/Final inboxes",
    "Design-token burn-down under baseline",
    "English-only i18n scaffold (Arabic removed per policy)",
])
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(5), "Fit & skills bugfix", [
    "Compare skills: matched-first tags + exact +N overflow",
    "Smart Sourcing: same overflow pattern",
    "Compare no longer hardcodes 0% as Good Match",
    "Real fit score when a vacancy is selected",
    "Without vacancy: Select a position (not fake 0%)",
    "Branch grok/ui-density-responsive merged by business owner",
])

# 11 Deployment
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Environments & Deployment", "How the system is packaged")
bullets(s, Inches(0.8), Inches(1.4), Inches(7), Inches(5), [
    "Docker multi-stage build targets: api, web, worker",
    "compose.production.yml wires api + web + worker + documents volume",
    "Web container proxies API via API_UPSTREAM",
    "Env via RECRUITFLOW_ENV_FILE / Dokploy Environment tab",
    "deploy/ contains nginx template, Hostinger–Vercel notes, runtime validators",
    "Never commit .env — use .env.example as contract",
], size=16)
card(s, Inches(8.2), Inches(1.5), Inches(4.5), Inches(4.6), "Core runtime needs", [
    "Node 24 + pnpm 11 (local)",
    "PostgreSQL",
    "Redis (queues)",
    "Document storage path/volume",
    "Affinda API key (CV parse)",
    "WEB_ORIGIN + secrets",
])

# 12 Commands
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Essential Commands", "Day-1 operations cheat sheet")
lines = [
    ("Local setup", 16, True, NAVY),
    ("pnpm install", 14, False, SLATE),
    ("pnpm db:generate && pnpm db:migrate:deploy && pnpm db:seed", 14, False, SLATE),
    ("pnpm dev:api   |   pnpm dev:web   |   pnpm dev:worker", 14, False, SLATE),
    ("", 10, False, SLATE),
    ("Quality gates", 16, True, NAVY),
    ("pnpm typecheck   |   pnpm lint   |   pnpm test", 14, False, SLATE),
    ("pnpm check:design-tokens   |   pnpm test:security", 14, False, SLATE),
    ("", 10, False, SLATE),
    ("Database", 16, True, NAVY),
    ("pnpm db:migrate:status   |   pnpm db:migrate:deploy", 14, False, SLATE),
]
add_text(s, Inches(0.8), Inches(1.4), Inches(12), Inches(5.5), lines)

# 13 Deferrals
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Known Deferrals & Risks", "Be explicit with IT — avoid false expectations")
bullets(s, Inches(0.8), Inches(1.4), Inches(12), Inches(5.5), [
    "SSO / SAML / OAuth — not implemented (username/password JWT only)",
    "Autonomous AI hiring decisions — not in product; fit scoring is assistive",
    "Full production email SMTP / WhatsApp / SMS — limited / template-oriented",
    "Payroll / full HCM replacement — out of scope",
    "Legacy palette debt still being burned down (gate prevents growth)",
    "Document/CV binary storage hardening continues via workers & storage path",
    "Support tip: Compare fit requires a selected vacancy to show real %",
], size=16)

# 14 Handover checklist
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "IT Handover Checklist", "Recommended acceptance items")
left_items = [
    "Repo access (GitHub/GitLab) + branch protection",
    "Production/staging env inventory documented",
    "Secrets vault ownership transferred",
    "DB backup & restore drill scheduled",
    "Docker/compose deploy verified by IT",
    "TLS / domain / reverse proxy confirmed",
]
right_items = [
    "Admin user break-glass process defined",
    "Monitoring/alerting hooks (API health, worker)",
    "Runbook links in docs/operations reviewed",
    "Support triage matrix (app vs infra)",
    "Release cadence & rollback owner named",
    "Security review of CORS/JWT cookies done",
]
bullets(s, Inches(0.7), Inches(1.4), Inches(5.8), Inches(5), left_items, size=15)
bullets(s, Inches(6.8), Inches(1.4), Inches(5.8), Inches(5), right_items, size=15)

# 15 Next steps
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title_bar(s, "Recommended Next Steps", "30 / 60 / 90 day view")
card(s, Inches(0.5), Inches(1.5), Inches(4), Inches(4.6), "0–30 days", [
    "IT shadow deploy + backup restore",
    "Access & secrets transfer",
    "Support mailbox / escalation path",
    "Smoke test critical hiring paths",
])
card(s, Inches(4.7), Inches(1.5), Inches(4), Inches(4.6), "30–60 days", [
    "Monitoring dashboards live",
    "SSO discovery workshop (if required)",
    "Perf & backup SLO agreement",
    "Training for L1/L2 support",
])
card(s, Inches(8.9), Inches(1.5), Inches(4), Inches(4.6), "60–90 days", [
    "Hardening backlog prioritization",
    "DR exercise",
    "Capacity review (DB/Redis/storage)",
    "Release governance finalized",
])

# 16 Closing
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, Inches(5.8), W, Inches(1.7), GREEN)
add_text(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(2), [
    ("Thank you", 44, True, WHITE),
    ("Questions from IT Department welcome", 22, False, RGBColor(0xB8, 0xD4, 0xEB)),
    ("RecruitFlow — ready for operational ownership transfer", 16, False, RGBColor(0xB8, 0xD4, 0xEB)),
])
add_text(s, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.8), [
    ("Artifacts: docs/architecture · docs/operations · deploy/ · RELEASE_NOTES.md · this deck in docs/handover/", 14, False, WHITE),
])

out = Path(r"D:\Projects\Recruitment Workflow System\docs\handover\RecruitFlow_IT_Handover_Presentation.pptx")
prs.save(out)
print(f"Wrote {out}")
print(f"Slides: {len(prs.slides)}")


